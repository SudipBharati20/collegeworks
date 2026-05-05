import urllib.request
import urllib.parse
from docx import Document
from docx.shared import Inches
import re
import os

def update_document():
    # 1. Download Flowchart
    print("Downloading flowchart...")
    dot = """digraph G {
        node [shape=box, style=filled, fillcolor=lightblue, fontname="Arial"];
        Start [shape=ellipse, fillcolor=lightgreen];
        Init [label="Initialize System\\n(Load Medicines)"];
        Menu [label="Display Menu"];
        Input [label="Get User Choice", shape=parallelogram, fillcolor=lightyellow];
        Cond1 [label="Choice == 1?", shape=diamond, fillcolor=orange];
        Op1 [label="Display Medicines"];
        Cond2 [label="Choice == 2?", shape=diamond, fillcolor=orange];
        Op2 [label="Process Sales"];
        Cond3 [label="Choice == 3?", shape=diamond, fillcolor=orange];
        Op3 [label="Process Restock"];
        Cond4 [label="Choice == 4?", shape=diamond, fillcolor=orange];
        Invalid [label="Print Error"];
        End [shape=ellipse, fillcolor=lightpink];

        Start -> Init;
        Init -> Menu;
        Menu -> Input;
        Input -> Cond1;
        Cond1 -> Op1 [label="Yes"];
        Op1 -> Menu;
        Cond1 -> Cond2 [label="No"];
        Cond2 -> Op2 [label="Yes"];
        Op2 -> Menu;
        Cond2 -> Cond3 [label="No"];
        Cond3 -> Op3 [label="Yes"];
        Op3 -> Menu;
        Cond3 -> Cond4 [label="No"];
        Cond4 -> End [label="Yes"];
        Cond4 -> Invalid [label="No"];
        Invalid -> Menu;
    }"""

    encoded_dot = urllib.parse.quote(dot)
    url = f"https://quickchart.io/graphviz?format=png&graph={encoded_dot}"
    img_path = "flowchart.png"
    urllib.request.urlretrieve(url, img_path)

    # 2. Modify docx
    print("Modifying focxxx.docx...")
    doc = Document("focxxx.docx")

    doc.add_heading('System Flowchart', level=1)
    doc.add_picture(img_path, width=Inches(5.5))

    doc.add_heading('Pseudocode', level=1)

    pseudocode = [
        "START",
        "  INITIALIZE MedStore System",
        "  LOAD medicines FROM file",
        "",
        "  WHILE True DO",
        "    DISPLAY \"MEDSTORE WHOLESALE MENU\"",
        "    DISPLAY \"1. View Medicine Inventory\"",
        "    DISPLAY \"2. Process Sales Transaction\"",
        "    DISPLAY \"3. Process Restock Transaction\"",
        "    DISPLAY \"4. Exit\"",
        "",
        "    PROMPT user FOR choice",
        "",
        "    IF choice EQUALS 1 THEN",
        "      CALL display_medicines",
        "    ELSE IF choice EQUALS 2 THEN",
        "      CALL handle_sales",
        "    ELSE IF choice EQUALS 3 THEN",
        "      CALL handle_restock",
        "    ELSE IF choice EQUALS 4 THEN",
        "      PRINT \"Exiting MedStore System\"",
        "      BREAK",
        "    ELSE",
        "      PRINT \"Error: Invalid choice\"",
        "    END IF",
        "  END WHILE",
        "END"
    ]

    for line in pseudocode:
        p = doc.add_paragraph()
        # Ensure courier font for pseudocode
        p.style.font.name = 'Courier New'
        # Split the line keeping whitespace and punctuation
        tokens = re.split(r'(\s+|[^a-zA-Z0-9\s]+)', line)
        for token in tokens:
            if not token:
                continue
            run = p.add_run(token)
            # Check if the token is entirely uppercase letters
            if token.isupper() and any(c.isalpha() for c in token):
                run.bold = True
            run.font.name = 'Courier New'

    doc.save("focxxx.docx")
    print("Successfully updated focxxx.docx with flowchart and pseudocode.")

if __name__ == "__main__":
    update_document()
